import sys,json,html,re,hashlib
from pathlib import Path
sys.path.insert(0,'/tmp/mat_reco_manual_deps')
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from blueprint import *
from steps_content import STEPS
from validation_data import *
NAME='Manuel-utilisateur-global-mat_reco-v0.2';DATE='14 octobre 2026'
CAPS=json.loads((OUT/'edition/Index-captures-v0.2.json').read_text())
ORDER={'F':0,'C':1,'A':2,'S':3,'E':4}
for i,c in enumerate(sorted(CAPS,key=lambda r:(ORDER[r['id'][0]],r['id'])),1):c['figure']=i
BYCODE={}
for c in CAPS:BYCODE.setdefault(c['etape'],[]).append(c)
REUSE={'F01':'F09','F03':'F04','C01':'C02','C08':'A26','C09':'S14','A01':'C07','A28':'A26'}
GRAPH_MAP={'F':'G02_flux_fournisseur','C':'G03_flux_client','A':'G04_flux_atelier_stock','S':'G06_relations_documents'}
ICON=KIT/'07_ICONES'
class Book:
 def __init__(self,title,name):
  self.name=name;self.title=title;self.doc=Document();self.parts=[];self.buf=[];self.sections=[];self.table_no=0;self.bid=0
  s=self.doc.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=s.bottom_margin=Cm(1.3);s.left_margin=s.right_margin=Cm(1.6);s.header_distance=s.footer_distance=Cm(.6)
  for style in ['Normal','List Bullet','List Number']:
   st=self.doc.styles[style];st.font.name='Calibri';st.font.size=Pt(10);st.paragraph_format.space_after=Pt(4);st.paragraph_format.line_spacing=1.04
  for style,size in [('Title',30),('Heading 1',18),('Heading 2',12)]:
   st=self.doc.styles[style];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('174C72');st.paragraph_format.space_before=Pt(5);st.paragraph_format.space_after=Pt(6)
  self.doc.styles['Caption'].font.size=Pt(8);self.doc.styles['Caption'].font.color.rgb=RGBColor.from_string('617487');self.doc.styles['Caption'].paragraph_format.space_after=Pt(3)
  self.doc.sections[0].header.paragraphs[0].text='mat_reco  |  Guide global des parcours  |  V0.2 — À valider'
  self.doc.sections[0].header.paragraphs[0].style='Caption'
  p=self.doc.sections[0].footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.add_run(DATE+'   •   Page ').font.size=Pt(8);f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
  self.doc.core_properties.title=title;self.doc.core_properties.subject='Fournisseur · Client · Atelier · Stock';self.doc.core_properties.author='Base éditoriale mat_reco pour validation documentaire et Skywork'
 def close(self):
  if self.buf:self.parts.append('<section id="'+html.escape(self.sections[-1][0])+'">'+''.join(self.buf)+'</section>');self.buf=[]
 def section(self,id,title,kicker=''):
  self.close()
  if self.sections:self.doc.add_page_break()
  self.sections.append((id,title));p=self.doc.add_heading(title,1);self.bid+=1
  b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(self.bid));b.set(qn('w:name'),id);p._p.insert(0,b);e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(self.bid));p._p.append(e)
  self.buf.append(f'<h1>{html.escape(title)}</h1>')
  if kicker:self.p(kicker,'caption')
 def p(self,text,kind='p'):
  p=self.doc.add_paragraph(text,style='Caption' if kind=='caption' else None)
  self.buf.append('<p class="'+kind+'">'+html.escape(text)+'</p>')
 def label(self,icon,label,text,color='blue'):
  p=self.doc.add_paragraph();p.paragraph_format.space_after=Pt(3)
  image_path=ICON/(icon+'.png')
  if image_path.exists():p.add_run().add_picture(str(image_path),width=Cm(.33))
  r=p.add_run('  '+label+'  ');r.bold=True;r.font.color.rgb=RGBColor.from_string({'blue':'246FA9','green':'248767','orange':'A97118','red':'B94646','violet':'7355AA'}[color]);p.add_run(text)
  rel=image_path.relative_to(OUT).as_posix();self.buf.append(f'<p class="label {color}"><img src="{rel}"><b>{html.escape(label)}</b> {html.escape(text)}</p>')
 def box(self,label,text,color='orange'):
  self.label('attention' if color=='orange' else 'resultat',label,text,color)
  pp=self.doc.paragraphs[-1]._p.get_or_add_pPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'FFF3D9' if color=='orange' else 'EAF7F0');pp.append(sh)
  self.buf[-1]=self.buf[-1].replace('<p class="label','<p class="box label')
 def pic(self,path,caption='',max_h=7.4,width=17.2):
  path=Path(path);assert path.exists(),path
  with Image.open(path) as im:w=min(width,max_h*im.width/im.height)
  p=self.doc.add_paragraph();p.paragraph_format.keep_with_next=bool(caption);p.paragraph_format.space_after=Pt(3);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(path),width=Cm(w))
  if caption:self.doc.add_paragraph(caption,style='Caption')
  rel=path.relative_to(OUT).as_posix();self.buf.append(f'<figure><img src="{rel}" style="width:{w:.3f}cm"><figcaption>{html.escape(caption)}</figcaption></figure>')
 def list(self,rows,ordered=False):
  for i,t in enumerate(rows,1):self.doc.add_paragraph((str(i)+'. ' if ordered else '• ')+t)
  tag='ol' if ordered else 'ul';self.buf.append('<'+tag+'>'+''.join('<li>'+html.escape(t)+'</li>' for t in rows)+'</'+tag+'>')
 def table(self,headers,rows,title=''):
  self.table_no+=1;label=f'Tableau {self.table_no}'+(' — '+title if title else '');self.p(label,'caption')
  tab=self.doc.add_table(rows=1,cols=len(headers));tab.style='Table Grid';hdr=tab.rows[0];hdr._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
  for c,t in zip(hdr.cells,headers):
   c.text=t;sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'174C72');c._tc.get_or_add_tcPr().append(sh)
   for r in c.paragraphs[0].runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255);r.font.size=Pt(9)
  for i,row in enumerate(rows):
   tr=tab.add_row();tr._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
   for c,t in zip(tr.cells,row):
    c.text=str(t)
    if i%2==0:sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F1F5F9');c._tc.get_or_add_tcPr().append(sh)
    for p in c.paragraphs:
     p.paragraph_format.space_after=Pt(2)
     for r in p.runs:r.font.size=Pt(8.7)
  self.buf.append('<table><thead><tr>'+''.join('<th>'+html.escape(t)+'</th>' for t in headers)+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+html.escape(str(x))+'</td>' for x in row)+'</tr>' for row in rows)+'</tbody></table>')
 def action_controls(self,actions,controls):
  table=self.doc.add_table(rows=1,cols=2);table.autofit=False;table.columns[0].width=Cm(9.5);table.columns[1].width=Cm(8.1)
  for c,label,icon,rows in [(table.cell(0,0),'Actions à réaliser','action',actions),(table.cell(0,1),'Points de contrôle','controle',controls)]:
   p=c.paragraphs[0];p.add_run().add_picture(str(ICON/(icon+'.png')),width=Cm(.33));r=p.add_run(' '+label);r.bold=True;r.font.color.rgb=RGBColor.from_string('246FA9' if icon=='action' else 'A97118')
   for i,t in enumerate(rows,1):
    p=c.add_paragraph((str(i)+'. ' if icon=='action' else '□ ')+t);p.paragraph_format.space_after=Pt(3)
    for r in p.runs:r.font.size=Pt(9.7)
  self.buf.append('<div class="actions-controls"><div><h2>Actions à réaliser</h2><ol>'+''.join('<li>'+html.escape(x)+'</li>' for x in actions)+'</ol></div><div><h2>Points de contrôle</h2><ul>'+''.join('<li>'+html.escape(x)+'</li>' for x in controls)+'</ul></div></div>')
 def save(self):
  self.close();self.doc.save(OUT/(self.name+'.docx'))
  css='''@page{size:A4;margin:13mm 16mm 14mm}*{box-sizing:border-box}body{font:10pt/1.13 Arial,sans-serif;color:#263c4d;margin:0}section{break-before:page}section:first-child{break-before:auto}h1{font-size:18pt;color:#174c72;margin:0 0 6pt}h2{font-size:11.5pt;margin:7pt 0 4pt;color:#246fa9}p{margin:0 0 4pt}.caption,figcaption{font-size:8pt;color:#617487}.caption{margin-bottom:6pt}.label img{width:3.3mm;height:3.3mm;vertical-align:middle;margin-right:3pt}.label b{margin-right:3pt}.blue b{color:#246fa9}.green b{color:#248767}.orange b{color:#a97118}.red b{color:#b94646}.violet b{color:#7355aa}.box{padding:5pt 7pt;background:#fff3d9;border-left:2pt solid #c38a26}.box.green{background:#eaf7f0;border-color:#248767}figure{margin:5pt 0;text-align:center;break-inside:avoid}figure img{max-width:100%;height:auto;display:block;margin:auto}figcaption{text-align:left;margin-top:3pt}ol,ul{margin:4pt 0 6pt;padding-left:17pt}li{margin:0 0 3pt}.actions-controls{display:grid;grid-template-columns:1.12fr 1fr;gap:15pt;break-inside:avoid}table{width:100%;border-collapse:collapse;font-size:8.5pt;margin:5pt 0 8pt}th,td{padding:5pt;border:1px solid #cfdae3;text-align:left;vertical-align:top}th{background:#174c72;color:#fff}tr{break-inside:avoid}thead{display:table-header-group}tbody tr:nth-child(odd){background:#f1f5f9}@media screen{body{background:#e8eff4}section{margin:20px auto;background:white;padding:13mm 16mm;max-width:210mm;min-height:297mm;box-shadow:0 1px 8px #b6c8d5}}'''
  (OUT/(self.name+'.html')).write_text('<!doctype html><html lang="fr"><meta charset="utf-8"><title>'+html.escape(self.title)+'</title><style>'+css+'</style><body>'+''.join(self.parts)+'</body></html>')
  return self.sections
