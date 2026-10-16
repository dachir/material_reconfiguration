import sys,json,hashlib,shutil
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from collections import defaultdict,Counter
from lxml import etree
sys.path.insert(0,'/tmp/mat_reco_manual_deps')
from pypdf import PdfReader
from openpyxl import load_workbook
from blueprint import *
from steps_content import STEPS
# Canonicalise identical images so the library is a reusable asset collection.
index_path=OUT/'edition/Index-captures-v0.2.json'
rows=json.loads(index_path.read_text());seen={};removed=[]
for r in rows:
 p=KIT/r['chemin'];sha=hashlib.sha256(p.read_bytes()).hexdigest();r['sha256_asset']=sha
 if sha in seen:
  original=seen[sha];old=p;r['asset_original_id']=original['id'];r['chemin']=original['chemin'];r['nom_fichier']=original['nom_fichier'];r['annotations']=r['annotations'] if ('Visuel partagé avec '+original['id']) in r['annotations'] else (r['annotations']+'; ' if r['annotations'] else '')+'Visuel partagé avec '+original['id']
  if old!=KIT/r['chemin'] and old.exists():old.unlink();removed.append(old.name)
 else:seen[sha]=r;r['asset_original_id']=r['id']
index_path.write_text(json.dumps(rows,ensure_ascii=False,indent=2))
# Keep all 60 figure usages in the index, backed by 54 different image files.
from build_indexes import *
readme='''MANUEL UTILISATEUR GLOBAL MAT_RECO — V0.2
BASE ÉDITORIALE À VALIDER — PRÉPARATION SKYWORK

Commencer par Manuel-utilisateur-global-mat_reco-v0.2.docx ou sa version PDF.
Le guide contient 64 fiches métier, 12 exceptions et 3 exercices guidés.

Contenu principal :
- Manuel Word modifiable et PDF de lecture (104 pages).
- Points-a-valider : 16 décisions ouvertes, limites et preuves attendues.
- Carte-des-flux : les 7 grands schémas en document Word / PDF.
- MANUEL_GLOBAL_MAT_RECO : captures par flux, graphes PNG/SVG, 13 icônes,
  storyboard et 64 textes de fiches séparés.
- INDEX_CAPTURES.xlsx : 60 usages de figures, 54 images distinctes,
  liens relatifs vers chaque image et 5 feuilles de contrôle.
- édition : données éditoriales structurées, jeu de démonstration et plan.
- validation : audit, sources, manifest et contrôles documentaires.

PROVENANCE ET LIMITES
Les interfaces atelier, services et dialogue Repack utilisent des composants
réels exécutés en démonstration avec données / réponses simulées en mémoire.
Les autres formulaires sont des illustrations reconstituées, identifiées.
Aucune capture en session métier authentifiée n’est revendiquée.
Aucun mouvement métier n’a été soumis pour produire ce pack.
Les 60 usages indexés regroupent 26 usages de démonstration et 34 reconstitutions ;
les images identiques ont été partagées pour éviter les fichiers redondants.

L’ordre de soumission MCP / Repack et les chemins de facture restent à valider.
Ne pas transformer ces décisions ouvertes en instructions officielles.
Le PDF est une version de lecture Chromium du même contenu que le Word,
avec une pagination susceptible de différer d’un rendu Microsoft Word.

POUR SKYWORK
Conserver codes de chapitres, numéros de figure, messages pédagogiques,
mentions de provenance et liens C07 → A01 / S14 → C08.
Les données sont fictives ; les variantes de scène sont signalées.
Extraire tout le ZIP avant d’ouvrir les liens relatifs du classeur.
'''
(OUT/'00_LISEZ_MOI.txt').write_text(readme)
(KIT/'08_STORYBOARD/BRIEF_SKYWORK.txt').write_text(readme)
(OUT/'edition/Comparaison-v0.1-v0.2.md').write_text('''# Écarts de la v0.2

- Trois flux métier convergents, à la place d’une succession de fonctions surtout MCP.
- 64 fiches stables F/C/A/S avec position, objectif, raison, actions, contrôles, résultat, transition.
- Circuit fournisseur complet : PO conditionnel, PR ou Receipt Plan, séries, disponibilité.
- Circuit client complet : devis, services, commande, passage atelier, livraison et facture.
- Repack découpe et Repack recomposition séparés.
- 7 grands graphes, 64 mini-graphes et 13 icônes éditoriales.
- 60 usages de captures indexés ; 54 images distinctes après mutualisation.
- 12 exceptions reliées aux fiches d’origine et 3 exercices de bout en bout.
- Index Excel, storyboard et textes séparés pour contrôle documentaire / Skywork.
- Limites explicites : aucune capture en session métier authentifiée, circuits et activation site à confirmer.
''')
# Include reproducible editorial sources, not business app source or site secrets.
src=OUT/'edition/sources_generateur';src.mkdir(exist_ok=True)
for p in (ROOT/'docs/user_manual_v02').glob('*'):
 if p.is_file() and p.suffix in ['.py','.cjs','.md']:shutil.copy2(p,src/p.name)
checks=[]
for p in OUT.glob('*.docx'):
 with ZipFile(p) as z:
  assert z.testzip() is None
  for name in z.namelist():
   if name.endswith('.xml'):etree.fromstring(z.read(name))
  xml=etree.fromstring(z.read('word/document.xml'));ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
  ids=xml.xpath('//w:bookmarkStart/@w:id',namespaces=ns);assert len(ids)==len(set(ids))
  if p.name.startswith('Manuel'):
   anchors=xml.xpath('//w:bookmarkStart/@w:name',namespaces=ns);assert set(STEPS).issubset(set(anchors))
  checks.append({'fichier':p.name,'controle':'ZIP intègre, XML valides, signets uniques','resultat':'OK'})
for p in OUT.glob('*.pdf'):
 r=PdfReader(p);assert all(len(page.extract_text().strip())>100 for page in r.pages)
 checks.append({'fichier':p.name,'pages':len(r.pages),'controle':'Lecture PDF et absence de page quasi vide','resultat':'OK'})
for r in rows:assert (KIT/r['chemin']).exists()
wb=load_workbook(OUT/'INDEX_CAPTURES.xlsx');assert wb['Captures'].max_row==61
assert len(set(r['id'] for r in rows))==60 and len(set(r['chemin'] for r in rows))==54
assert all((KIT/'06_FLOWCHARTS'/('position_'+code+'.png')).exists() for code in STEPS)
checks.append({'controle':'60 lignes indexées, 54 images distinctes, 64 repères, 64 fiches présentes','resultat':'OK'})
(OUT/'validation/Controle-documentaire-v0.2.json').write_text(json.dumps({'version':'0.2','checks':checks,'mutualisations':removed,'limites':'Validation documentaire locale, pas une recette de transactions métier ni un rendu Word contrôlé dans Microsoft Word.'},ensure_ascii=False,indent=2))
# Single zip includes every final artefact, independent of server paths.
zip_path=OUT/'MANUEL_GLOBAL_MAT_RECO_v0.2_Tous_les_assets.zip'
with ZipFile(zip_path,'w',ZIP_DEFLATED,compresslevel=6) as z:
 for p in sorted(OUT.rglob('*')):
  if p.is_file() and p.suffix!='.zip' and p.name not in ['index.html','SHA256SUMS.txt'] and '__pycache__' not in p.parts:
   z.write(p,'MANUEL_GLOBAL_MAT_RECO_v0.2/'+p.relative_to(OUT).as_posix())
with ZipFile(zip_path) as z:assert z.testzip() is None;count=len(z.namelist())
sha=hashlib.sha256(zip_path.read_bytes()).hexdigest();(OUT/'SHA256SUMS.txt').write_text(sha+'  '+zip_path.name+'\n')
(OUT/'index.html').write_text('''<!doctype html><html lang="fr"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>mat_reco — Pack global v0.2</title><style>body{font:18px Arial;background:#eef3f7;color:#23394c;padding:40px}main{max-width:850px;background:white;padding:35px;border-radius:12px;margin:auto}a{color:#246fa9}li{padding:10px}h1{color:#174c72}</style><main><h1>Manuel global mat_reco · V0.2</h1><p>Base éditoriale à valider ensemble, prête pour Skywork.</p><p><a download href="MANUEL_GLOBAL_MAT_RECO_v0.2_Tous_les_assets.zip"><b>Télécharger tous les assets dans un ZIP</b></a></p><ul><li><a download href="Manuel-utilisateur-global-mat_reco-v0.2.docx">Manuel Word</a></li><li><a href="Manuel-utilisateur-global-mat_reco-v0.2.pdf">Manuel PDF</a></li><li><a href="Points-a-valider-mat_reco-v0.2.pdf">Points à valider</a></li><li><a download href="INDEX_CAPTURES.xlsx">Index Excel</a></li><li><a href="Carte-des-flux-mat_reco-v0.2.pdf">Carte des flux</a></li></ul><p>64 fiches · 12 exceptions · 3 exercices · 54 images distinctes pour 60 usages · 7 grands graphes · 13 icônes.</p><p>Les captures en démonstration et les vues reconstituées sont identifiées ; aucune session métier authentifiée n’est présentée comme capturée.</p></main></html>''')
print('FINAL ZIP:',zip_path.name,'|',count,'files |',round(zip_path.stat().st_size/1024/1024,1),'MiB | SHA256',sha)
