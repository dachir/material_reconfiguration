import json,sys,hashlib,re,csv,shutil
from pathlib import Path
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from blueprint import *
from validation_data import *
from steps_content import STEPS
from doc_utils import Book
CAPS=json.loads((OUT/'edition/Index-captures-v0.2.json').read_text())
# Separate maps book, seven complete reusable diagrams.
flowbook=Book('Carte des flux — mat_reco v0.2','Carte-des-flux-mat_reco-v0.2')
flowbook.section('maps_cover','Carte des flux · mat_reco v0.2','Base de lecture et de validation des trois parcours convergents.')
flowbook.p('Les schémas servent de carte du manuel. Le fournisseur apporte la matière, le client exprime le besoin et le Flux 3 les rapproche. Le Repack de découpe et celui de recomposition sont distincts. Les responsabilités et l’ordre de soumission restent à approuver.')
graphs=[('G01','Carte globale des trois flux',KIT/'00_CARTE_GLOBALE/G01_carte_globale_3_flux.png'),('G02','Flux Fournisseur',KIT/'06_FLOWCHARTS/G02_flux_fournisseur.png'),('G03','Flux Client',KIT/'06_FLOWCHARTS/G03_flux_client.png'),('G04','Flux Atelier / Stock',KIT/'06_FLOWCHARTS/G04_flux_atelier_stock.png'),('G05','Compensation',KIT/'06_FLOWCHARTS/G05_compensation.png'),('G06','Relations entre documents',KIT/'06_FLOWCHARTS/G06_relations_documents.png'),('G07','Décision en cas de casse',KIT/'06_FLOWCHARTS/G07_decision_casse.png')]
for code,title,path in graphs:flowbook.section(code,code+' · '+title);flowbook.pic(path,'Schéma '+code+' · Proposition à valider ; fichier PNG et SVG séparés fournis.',23)
flowbook.save()
# Index workbook contains clickable paths to full-size images and reusable SVGs.
wb=Workbook();ws=wb.active;ws.title='Captures';headers=['ID capture','Figure manuel','Nom fichier','Flux','Étape','Écran / document','Message pédagogique','Statut','Provenance','Source de référence','Annotations','Validation site','Chemin asset','Résolution']
ws.append(headers)
for c in sorted(CAPS,key=lambda c:(c['flux'],c['id'])):
 ws.append([c['id'],c['figure'],c['nom_fichier'],c['flux'],c['etape'],c['ecran'],c['message'],c['statut'],c['provenance'],c['source'],c['annotations'],c['validation_site'],c['chemin'],c['resolution']]);cell=ws.cell(ws.max_row,13);cell.hyperlink=c['chemin'];cell.style='Hyperlink'
ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
for j,w in enumerate([14,12,46,26,12,32,66,20,70,72,52,43,60,20],1):ws.column_dimensions[get_column_letter(j)].width=w
ws2=wb.create_sheet('Graphes')
ws2.append(['ID','Carte','PNG','SVG','Usage'])
for code,t,p in graphs:
 rel=p.relative_to(KIT).as_posix();ws2.append([code,t,rel,rel.replace('.png','.svg'),'Proposition à valider ; conserver les décisions ouvertes.'])
 for col in [3,4]:ws2.cell(ws2.max_row,col).hyperlink=ws2.cell(ws2.max_row,col).value;ws2.cell(ws2.max_row,col).style='Hyperlink'
ws3=wb.create_sheet('Décisions ouvertes');ws3.append(['ID','Sujet','État actuel','Responsable','Validation attendue','Décision / preuve'])
for r in DECISIONS:ws3.append([*r,'À compléter'])
ws4=wb.create_sheet('Storyboard');ws4.append(['Séquence','Titre','Question utilisateur','Visuel dominant','Chapitres'])
for r in STORYBOARD:ws4.append(r)
ws5=wb.create_sheet('Lisez-moi');ws5.append(['Sujet','Règle'])
for r in [('Lot','Manuel v0.2 ; 64 fiches ; 60 usages de figures / 54 images distinctes ; 7 grands graphes ; 64 repères de position ; 13 icônes.'),('Provenance','26 démonstrations de composants réels ; 34 illustrations reconstituées ; aucune capture de session métier authentifiée.'),('Utilisation','Les liens relatifs fonctionnent après extraction complète du ZIP. Ouvrir les PNG en pleine résolution.'),('Skywork','Conserver les codes de fiche, numéros de figure et statuts de provenance.'),('Validation','L’ordre MCP / Repack et les variantes de facture restent ouverts ; ne pas les rendre obligatoires graphiquement.')]:ws5.append(r)
for sheet in wb:
 sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
 for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='174C72');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(vertical='center',wrap_text=True)
 sheet.row_dimensions[1].height=32
 for row in sheet.iter_rows(min_row=2):
  for cell in row:
   cell.alignment=Alignment(vertical='top',wrap_text=True)
   if cell.row%2==0:cell.fill=PatternFill('solid',fgColor='F1F5F9')
  sheet.row_dimensions[row[0].row].height=56 if sheet.title=='Captures' else 50
 if sheet!=ws:
  for col in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(col)].width=24 if col==1 else 64
wb.save(KIT/'INDEX_CAPTURES.xlsx');shutil.copy2(KIT/'INDEX_CAPTURES.xlsx',OUT/'INDEX_CAPTURES.xlsx')
with (KIT/'INDEX_CAPTURES.csv').open('w',newline='',encoding='utf-8-sig') as f:
 w=csv.writer(f,delimiter=';');w.writerow(headers)
 for c in CAPS:w.writerow([c['id'],c['figure'],c['nom_fichier'],c['flux'],c['etape'],c['ecran'],c['message'],c['statut'],c['provenance'],c['source'],c['annotations'],c['validation_site'],c['chemin'],c['resolution']])
# Editorial package contains text separately, no screenshots hidden in a Word only.
text='# Storyboard Skywork — mat_reco v0.2\n\n'
for num,t,q,visual,chap in STORYBOARD:text+=f'## Séquence {num} — {t}\n\nQuestion : {q}\n\nVisuel dominant : {visual}\n\nChapitres : {chap}\n\n'
text+='## Règles de finalisation\n\n- Conserver les codes F/C/A/S/E, les numéros de figure et les transitions.\n- Garder les statuts Réelle / Démonstration / Reconstituée visibles.\n- Ne jamais renommer une illustration comme une capture ERPNext réelle.\n- Ne pas transformer l’ordre de soumission MCP / Repack en règle officielle.\n- Montrer fournisseur et client convergeant vers l’atelier.\n- Distinguer Repack découpe et recomposition, même si les deux documents s’appellent Stock Entry.\n- Garder les mini-graphes « Vous êtes ici » et les messages d’attention.\n- Les données de formation sont fictives ; les variantes L / stock sont annoncées.\n'
(KIT/'08_STORYBOARD/Storyboard-Skywork-v0.2.md').write_text(text)
(KIT/'08_STORYBOARD/Storyboard-Skywork-v0.2.json').write_text(json.dumps(STORYBOARD,ensure_ascii=False,indent=2))
fichedir=KIT/'08_STORYBOARD/FICHES';fichedir.mkdir(exist_ok=True)
for code,st in STEPS.items():
 sections=[('# '+code+' — '+st['title'],''),('Position dans le flux',GROUPS[code[0]][0]),('Objectif',st['objective']),('Pourquoi cette étape existe',st['why']),('Quand l’utiliser',st['when']),('Qui l’exécute',st['role']),('Document',st['document']),('Entrées / prérequis',st['inputs']),('Actions','\n'.join(f'{i}. {s}' for i,s in enumerate(st['actions'],1))),('Contrôles','\n'.join('- '+s for s in st['controls'])),('Résultat',st['result']),('Étape suivante',st['next']),('Attention',st['attention'] or 'Ne pas poursuivre si un contrôle reste incohérent.'),('Source documentaire','\n'.join(st['sources']))]
 (fichedir/(code+'.md')).write_text('\n\n'.join((head+'\n\n'+body) if i==0 else ('## '+head+'\n\n'+body) for i,(head,body) in enumerate(sections)))
# Source manifest uses hashes, not credentials, records or unpublished business data.
sources=[]
sourcepaths=set(s for st in STEPS.values() for s in st['sources'])|set(s for c in CAPS for s in c['source'].split(';') if s)
for s in sorted(sourcepaths):
 p=ROOT/s if not s.startswith('erpnext/') else ROOT.parent/'erpnext/erpnext'/s.removeprefix('erpnext/')
 if p.is_dir():continue
 if p.exists():sources.append({'reference':s,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'bytes':p.stat().st_size})
 else:sources.append({'reference':s,'status':'Référence de composant / chemin à préciser'})
manifest={'version':'0.2','date':'2026-10-14','statut':'Base éditoriale à valider','fiches_metier':64,'exceptions':12,'captures':dict(Counter(c['statut'] for c in CAPS)),'captures_session_metier_reelle':0,'graphes_principaux':7,'reperes_position':64,'icones':13,'sources':sources,'limites':['Pas de capture en session métier authentifiée.','Pas de soumission de mouvements métier pour produire le guide.','Procédures de soumission et chemins de facturation à valider.']}
(OUT/'validation/Manifest-documentaire-v0.2.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
(OUT/'validation/Points-a-valider-v0.2.md').write_text('# Points à valider\n\n'+'\n\n'.join(f'## {id} — {topic}\n\nÉtat : {state}\n\nResponsable : {owner}\n\nValidation : {test}' for id,topic,state,owner,test in DECISIONS))
print('Phase 5 : classeur index, cartes Word, storyboard, 64 fiches séparées, manifest et décisions générés.')
