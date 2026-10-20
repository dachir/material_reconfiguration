from pathlib import Path
import sys,json,hashlib,html
sys.path.insert(0,'/tmp/mat_reco_manual_deps')
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from manual_content import TITLE,SUBTITLE,VERSION,DATE,PAGES
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'mat_reco/public/manuals/mat_reco-v0-1';IMG=OUT/'images'
NAME='Manuel-utilisateur-mat_reco-v0-1'
BLUE='174C72';GRAY='516578'
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(1.6);sec.left_margin=sec.right_margin=Cm(1.7)
sec.header_distance=sec.footer_distance=Cm(.7)
for style_name in ['Normal','List Bullet','List Number']:
 st=doc.styles[style_name];st.font.name='Calibri';st.font.size=Pt(10.5);st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.08
for name,size in [('Title',34),('Subtitle',17),('Heading 1',20),('Heading 2',13)]:
 st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(BLUE)
 st.paragraph_format.space_after=Pt(8);st.paragraph_format.space_before=Pt(9)
doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].font.color.rgb=RGBColor.from_string(GRAY)
doc.styles['Caption'].paragraph_format.space_after=Pt(7)
head=sec.header.paragraphs[0];head.text='mat_reco  /  Manuel utilisateur  /  Version à valider';head.style=doc.styles['Caption']
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;r=f.add_run('V0.1 · '+DATE+'   |   Page ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
props=doc.core_properties;props.title=TITLE;props.subject=SUBTITLE;props.author='Manuel préparé pour validation utilisateur';props.keywords='mat_reco, MCP, découpe, stock, Repack, Retour Terrain';props.comments='Source : code local analysé. Données des captures fictives.'
settings=OxmlElement('w:updateFields');settings.set(qn('w:val'),'true');doc.settings.element.append(settings)

def shade(cell,color):
 tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement('w:shd');e.set(qn('w:fill'),color);tcPr.append(e)
def no_split(row):
 pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
def bookmark(par,name):
 e=OxmlElement('w:bookmarkStart');e.set(qn('w:id'),str(int(name.rsplit('_',1)[1])));e.set(qn('w:name'),name);par._p.insert(0,e)
 end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(int(name.rsplit('_',1)[1])));par._p.append(end)
def link(par,label,anchor):
 e=OxmlElement('w:hyperlink');e.set(qn('w:anchor'),anchor);r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),BLUE);pr.append(c);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);e.append(r);par._p.append(e)

doc.add_paragraph('MATIÈRE  •  DÉCOUPE  •  STOCK',style='Subtitle')
doc.add_paragraph('\n')
doc.add_paragraph('mat_reco',style='Title')
doc.add_paragraph('Manuel utilisateur',style='Title')
doc.add_paragraph(SUBTITLE,style='Subtitle')
doc.add_paragraph(VERSION,style='Heading 2');doc.add_paragraph(DATE)
doc.add_picture(str(IMG/'01-atelier.png'),width=Cm(17.2))
doc.add_paragraph('Interface réelle · données de démonstration',style='Caption')
doc.add_paragraph('De la réception de matière au contrôle du Repack, avec incidents, rotation et compensation des découpes.')
doc.add_paragraph('À valider avec l’atelier et le responsable de stock avant diffusion officielle.')
doc.add_page_break();doc.add_heading('Sommaire',1)
for i,p in enumerate(PAGES):
 para=doc.add_paragraph();para.paragraph_format.space_after=Pt(3);link(para,p['title'],f'section_{i+1}')

sections=[]
for idx,page in enumerate(PAGES):
 doc.add_page_break();hp=doc.add_heading(page['title'],1);bookmark(hp,f'section_{idx+1}')
 frag=[f'<h1>{html.escape(page["title"])}</h1>']
 for item in page['items']:
  kind=item[0]
  if kind=='p':doc.add_paragraph(item[1]);frag.append('<p>'+html.escape(item[1])+'</p>')
  elif kind=='h2':doc.add_heading(item[1],2);frag.append('<h2>'+html.escape(item[1])+'</h2>')
  elif kind in ['steps','bullets']:
   number_id=None
   if kind=='steps':
    numbering=doc.part.numbering_part.element
    source_id=doc.styles['List Number'].element.pPr.numPr.numId.val
    abstract_id=numbering.num_having_numId(source_id).abstractNumId.val
    num=numbering.add_num(abstract_id);num.add_lvlOverride(ilvl=0).add_startOverride(1)
    number_id=num.numId
   for text in item[1]:
    para=doc.add_paragraph(text,style='List Number' if kind=='steps' else 'List Bullet')
    if number_id is not None:para._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val=number_id
   tag='ol' if kind=='steps' else 'ul';frag.append('<'+tag+'>'+''.join('<li>'+html.escape(t)+'</li>' for t in item[1])+'</'+tag+'>')
  elif kind=='note':
   para=doc.add_paragraph(item[1]);pr=para._p.get_or_add_pPr();s=OxmlElement('w:shd');s.set(qn('w:fill'),'FFF4D9');pr.append(s)
   para.paragraph_format.space_before=Pt(6);para.paragraph_format.space_after=Pt(8)
   for r in para.runs:r.font.size=Pt(10)
   frag.append('<aside>'+html.escape(item[1])+'</aside>')
  elif kind=='table':
   tab=doc.add_table(rows=1,cols=len(item[1]));tab.style='Table Grid'
   hdr=tab.rows[0];repeat=OxmlElement('w:tblHeader');hdr._tr.get_or_add_trPr().append(repeat)
   for c,text in zip(hdr.cells,item[1]):
    c.text=text;shade(c,BLUE)
    for r in c.paragraphs[0].runs:r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255);r.font.size=Pt(9.5)
   for j,row in enumerate(item[2]):
    tr=tab.add_row();no_split(tr)
    for c,text in zip(tr.cells,row):
     c.text=text
     if j%2==0:shade(c,'F1F5F9')
     for para in c.paragraphs:
      para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(2)
      for r in para.runs:r.font.size=Pt(9.3)
   doc.add_paragraph().paragraph_format.space_after=Pt(0)
   frag.append('<table><thead><tr>'+''.join('<th>'+html.escape(x)+'</th>' for x in item[1])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+html.escape(x)+'</td>' for x in row)+'</tr>' for row in item[2])+'</tbody></table>')
  elif kind=='fig':
   fp=(IMG/item[1]).resolve();assert fp.is_file(),fp
   with Image.open(fp) as im:assert im.width>100 and im.height>100
   para=doc.add_paragraph();para.paragraph_format.keep_with_next=True;para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.add_run().add_picture(str(fp),width=Cm(item[3]));doc.add_paragraph(item[2],style='Caption')
   rel=fp.relative_to(OUT).as_posix();frag.append(f'<figure><img style="width:{item[3]}cm" src="{rel}"><figcaption>{html.escape(item[2])}</figcaption></figure>')
 sections.append('<section id="section_'+str(idx+1)+'">'+''.join(frag)+'</section>')
doc.save(OUT/(NAME+'.docx'))
cover=f'<section class="cover"><p class="eyebrow">MATIÈRE · DÉCOUPE · STOCK</p><h1>mat_reco<br>Manuel utilisateur</h1><h2>{SUBTITLE}</h2><p><b>{VERSION}</b><br>{DATE}</p><figure><img src="images/01-atelier.png"><figcaption>Interface réelle · données de démonstration</figcaption></figure><p>De la réception de matière au contrôle du Repack, avec incidents, rotation et compensation des découpes.</p><aside>À valider avec l’atelier et le responsable de stock avant diffusion officielle.</aside></section>'
toc='<section class="toc"><h1>Sommaire</h1><ol>'+''.join(f'<li><a href="#section_{i+1}">{html.escape(p["title"])}</a></li>' for i,p in enumerate(PAGES))+'</ol></section>'
css='''@page{size:A4;margin:16mm 17mm 17mm}*{box-sizing:border-box}body{margin:0;color:#25364a;font-family:Arial,sans-serif;font-size:10.4pt;line-height:1.22}section{break-before:page}section:first-child{break-before:auto}h1{font-size:20pt;color:#174c72;margin:0 0 12pt}h2{font-size:13pt;color:#174c72;margin:12pt 0 7pt}p{margin:0 0 7pt}li{margin-bottom:5pt}ol,ul{margin:6pt 0 10pt;padding-left:22pt}table{width:100%;border-collapse:collapse;font-size:9.2pt;margin:8pt 0 12pt}th{background:#174c72;color:#fff}th,td{border:1px solid #ced8e1;padding:6pt;text-align:left;vertical-align:top}tbody tr:nth-child(odd){background:#f1f5f9}tr,figure,aside{break-inside:avoid}thead{display:table-header-group}figure{margin:10pt 0;text-align:center}figure img{max-width:100%;height:auto;display:block;margin:auto}figcaption{font-size:8.5pt;color:#516578;text-align:left;margin-top:5pt}aside{background:#fff4d9;border-left:3px solid #d1a038;padding:9pt;margin:9pt 0;font-size:9.6pt}.cover h1{font-size:36pt;margin-top:35pt}.cover h2{font-size:17pt}.cover figure{margin:24pt 0}.eyebrow{letter-spacing:1.7px;color:#174c72}.toc li{font-size:10.3pt;margin-bottom:5pt}.toc ol{list-style:none;padding:0}a{color:#174c72;text-decoration:none}@media screen{body{background:#dde5ed}section{background:white;max-width:210mm;min-height:297mm;padding:16mm 17mm;margin:25px auto;box-shadow:0 2px 9px #aabbcc}}'''
(OUT/(NAME+'.html')).write_text('<!doctype html><html lang="fr"><meta charset="utf-8"><title>'+TITLE+'</title><style>'+css+'</style><body>'+cover+toc+''.join(sections)+'</body></html>')
# Source coverage manifest: no secrets or business data, only paths / hashes of the app.
files=[]
for group in ['mat_reco','frontend/mcp-ux20-vue/src']:
 for p in sorted((ROOT/group).rglob('*')):
  if not p.is_file() or p.suffix not in ['.py','.js','.vue','.json','.css']:continue
  if any(x in p.parts for x in ['__pycache__','node_modules','mcp_vue_dist','manuals']):continue
  if p.suffix=='.js' and 'public' in p.parts and p.name.endswith('.bundle.js'):continue
  files.append({'path':p.relative_to(ROOT).as_posix(),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'bytes':p.stat().st_size})
manifest={'version':'0.1','date':DATE,'audience':SUBTITLE,'status':'À valider','scope':'Analyse locale des DocTypes, services, scripts, hooks, fixtures et atelier Vue. Captures de démonstration ; aucun mouvement métier soumis.','source_files':files,'chapters':len(PAGES),'figures':14,'flowcharts':2}
(OUT/'Analyse-et-sources-mat_reco-v0-1.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print('Created',NAME+'.docx','and HTML reading version;',len(PAGES),'chapters;',len(files),'source files in coverage manifest.')
