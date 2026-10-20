import sys,json,html,re,hashlib
from pathlib import Path
sys.path.insert(0,'/tmp/mat_reco_manual_deps')
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from blueprint import *
from steps_content import STEPS
from validation_data import *
NAME='Manuel-utilisateur-global-mat_reco-v0.2';DATE='14 octobre 2026'
CAPS=json.loads((OUT/'edition/Index-captures-v0.2.json').read_text())
ORDER={'F':0,'C':1,'A':2,'S':3,'E':4}
for i,c in enumerate(sorted(CAPS,key=lambda r:(ORDER[r['id'][0]],r['id'])),1):c['figure']=i
BYCODE={}
for c in CAPS:BYCODE.setdefault(c['etape'],[]).append(c)
REUSE={'F01':'F09','F03':'F04','C01':'C02','C08':'A26','C09':'S14','A01':'C07','A28':'A26'}
GRAPH_MAP={'F':'G02_flux_fournisseur','C':'G03_flux_client','A':'G04_flux_atelier_stock','S':'G06_relations_documents'}
ICON=KIT/'07_ICONES'
class Book:
 def __init__(self,title,name):
  self.name=name;self.title=title;self.doc=Document();self.parts=[];self.buf=[];self.sections=[];self.table_no=0;self.bid=0
  s=self.doc.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=s.bottom_margin=Cm(1.3);s.left_margin=s.right_margin=Cm(1.6);s.header_distance=s.footer_distance=Cm(.6)
  for style in ['Normal','List Bullet','List Number']:
   st=self.doc.styles[style];st.font.name='Calibri';st.font.size=Pt(10);st.paragraph_format.space_after=Pt(4);st.paragraph_format.line_spacing=1.04
  for style,size in [('Title',30),('Heading 1',18),('Heading 2',12)]:
   st=self.doc.styles[style];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('174C72');st.paragraph_format.space_before=Pt(5);st.paragraph_format.space_after=Pt(6)
  self.doc.styles['Caption'].font.size=Pt(8);self.doc.styles['Caption'].font.color.rgb=RGBColor.from_string('617487');self.doc.styles['Caption'].paragraph_format.space_after=Pt(3)
  self.doc.sections[0].header.paragraphs[0].text='mat_reco  |  Guide global des parcours  |  V0.2 — À valider'
  self.doc.sections[0].header.paragraphs[0].style='Caption'
  p=self.doc.sections[0].footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.add_run(DATE+'   •   Page ').font.size=Pt(8);f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
  self.doc.core_properties.title=title;self.doc.core_properties.subject='Fournisseur · Client · Atelier · Stock';self.doc.core_properties.author='Base éditoriale mat_reco pour validation documentaire et Skywork'
 def close(self):
  if self.buf:self.parts.append('<section id="'+html.escape(self.sections[-1][0])+'">'+''.join(self.buf)+'</section>');self.buf=[]
 def section(self,id,title,kicker=''):
  self.close()
  if self.sections:self.doc.add_page_break()
  self.sections.append((id,title));p=self.doc.add_heading(title,1);self.bid+=1
  b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(self.bid));b.set(qn('w:name'),id);p._p.insert(0,b);e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(self.bid));p._p.append(e)
  self.buf.append(f'<h1>{html.escape(title)}</h1>')
  if kicker:self.p(kicker,'caption')
 def p(self,text,kind='p'):
  p=self.doc.add_paragraph(text,style='Caption' if kind=='caption' else None)
  self.buf.append('<p class="'+kind+'">'+html.escape(text)+'</p>')
 def label(self,icon,label,text,color='blue'):
  p=self.doc.add_paragraph();p.paragraph_format.space_after=Pt(3)
  image_path=ICON/(icon+'.png')
  if image_path.exists():p.add_run().add_picture(str(image_path),width=Cm(.33))
  r=p.add_run('  '+label+'  ');r.bold=True;r.font.color.rgb=RGBColor.from_string({'blue':'246FA9','green':'248767','orange':'A97118','red':'B94646','violet':'7355AA'}[color]);p.add_run(text)
  rel=image_path.relative_to(OUT).as_posix();self.buf.append(f'<p class="label {color}"><img src="{rel}"><b>{html.escape(label)}</b> {html.escape(text)}</p>')
 def box(self,label,text,color='orange'):
  self.label('attention' if color=='orange' else 'resultat',label,text,color)
  pp=self.doc.paragraphs[-1]._p.get_or_add_pPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'FFF3D9' if color=='orange' else 'EAF7F0');pp.append(sh)
  self.buf[-1]=self.buf[-1].replace('<p class="label','<p class="box label')
 def pic(self,path,caption='',max_h=7.4,width=17.2):
  path=Path(path);assert path.exists(),path
  with Image.open(path) as im:w=min(width,max_h*im.width/im.height)
  p=self.doc.add_paragraph();p.paragraph_format.keep_with_next=bool(caption);p.paragraph_format.space_after=Pt(3);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(path),width=Cm(w))
  if caption:self.doc.add_paragraph(caption,style='Caption')
  rel=path.relative_to(OUT).as_posix();self.buf.append(f'<figure><img src="{rel}" style="width:{w:.3f}cm"><figcaption>{html.escape(caption)}</figcaption></figure>')
 def list(self,rows,ordered=False):
  for i,t in enumerate(rows,1):self.doc.add_paragraph((str(i)+'. ' if ordered else '• ')+t)
  tag='ol' if ordered else 'ul';self.buf.append('<'+tag+'>'+''.join('<li>'+html.escape(t)+'</li>' for t in rows)+'</'+tag+'>')
 def table(self,headers,rows,title=''):
  self.table_no+=1;label=f'Tableau {self.table_no}'+(' — '+title if title else '');self.p(label,'caption')
  tab=self.doc.add_table(rows=1,cols=len(headers));tab.style='Table Grid';hdr=tab.rows[0];hdr._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
  for c,t in zip(hdr.cells,headers):
   c.text=t;sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'174C72');c._tc.get_or_add_tcPr().append(sh)
   for r in c.paragraphs[0].runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255);r.font.size=Pt(9)
  for i,row in enumerate(rows):
   tr=tab.add_row();tr._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
   for c,t in zip(tr.cells,row):
    c.text=str(t)
    if i%2==0:sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F1F5F9');c._tc.get_or_add_tcPr().append(sh)
    for p in c.paragraphs:
     p.paragraph_format.space_after=Pt(2)
     for r in p.runs:r.font.size=Pt(8.7)
  self.buf.append('<table><thead><tr>'+''.join('<th>'+html.escape(t)+'</th>' for t in headers)+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+html.escape(str(x))+'</td>' for x in row)+'</tr>' for row in rows)+'</tbody></table>')
 def action_controls(self,actions,controls):
  table=self.doc.add_table(rows=1,cols=2);table.autofit=False;table.columns[0].width=Cm(9.5);table.columns[1].width=Cm(8.1)
  for c,label,icon,rows in [(table.cell(0,0),'Actions à réaliser','action',actions),(table.cell(0,1),'Points de contrôle','controle',controls)]:
   p=c.paragraphs[0];p.add_run().add_picture(str(ICON/(icon+'.png')),width=Cm(.33));r=p.add_run(' '+label);r.bold=True;r.font.color.rgb=RGBColor.from_string('246FA9' if icon=='action' else 'A97118')
   for i,t in enumerate(rows,1):
    p=c.add_paragraph((str(i)+'. ' if icon=='action' else '□ ')+t);p.paragraph_format.space_after=Pt(3)
    for r in p.runs:r.font.size=Pt(9.7)
  self.buf.append('<div class="actions-controls"><div><h2>Actions à réaliser</h2><ol>'+''.join('<li>'+html.escape(x)+'</li>' for x in actions)+'</ol></div><div><h2>Points de contrôle</h2><ul>'+''.join('<li>'+html.escape(x)+'</li>' for x in controls)+'</ul></div></div>')
 def save(self):
  self.close();self.doc.save(OUT/(self.name+'.docx'))
  css='''@page{size:A4;margin:13mm 16mm 14mm}*{box-sizing:border-box}body{font:10pt/1.13 Arial,sans-serif;color:#263c4d;margin:0}section{break-before:page}section:first-child{break-before:auto}h1{font-size:18pt;color:#174c72;margin:0 0 6pt}h2{font-size:11.5pt;margin:7pt 0 4pt;color:#246fa9}p{margin:0 0 4pt}.caption,figcaption{font-size:8pt;color:#617487}.caption{margin-bottom:6pt}.label img{width:3.3mm;height:3.3mm;vertical-align:middle;margin-right:3pt}.label b{margin-right:3pt}.blue b{color:#246fa9}.green b{color:#248767}.orange b{color:#a97118}.red b{color:#b94646}.violet b{color:#7355aa}.box{padding:5pt 7pt;background:#fff3d9;border-left:2pt solid #c38a26}.box.green{background:#eaf7f0;border-color:#248767}figure{margin:5pt 0;text-align:center;break-inside:avoid}figure img{max-width:100%;height:auto;display:block;margin:auto}figcaption{text-align:left;margin-top:3pt}ol,ul{margin:4pt 0 6pt;padding-left:17pt}li{margin:0 0 3pt}.actions-controls{display:grid;grid-template-columns:1.12fr 1fr;gap:15pt;break-inside:avoid}table{width:100%;border-collapse:collapse;font-size:8.5pt;margin:5pt 0 8pt}th,td{padding:5pt;border:1px solid #cfdae3;text-align:left;vertical-align:top}th{background:#174c72;color:#fff}tr{break-inside:avoid}thead{display:table-header-group}tbody tr:nth-child(odd){background:#f1f5f9}@media screen{body{background:#e8eff4}section{margin:20px auto;background:white;padding:13mm 16mm;max-width:210mm;min-height:297mm;box-shadow:0 1px 8px #b6c8d5}}'''
  (OUT/(self.name+'.html')).write_text('<!doctype html><html lang="fr"><meta charset="utf-8"><title>'+html.escape(self.title)+'</title><style>'+css+'</style><body>'+''.join(self.parts)+'</body></html>')
  return self.sections
b=Book('Manuel utilisateur global mat_reco',NAME)
b.section('cover','Manuel utilisateur global\nmat_reco','V0.2 · Base éditoriale pour Codex, contrôle documentaire et Skywork')
b.p('FOURNISSEUR  →  MATIÈRE\nCLIENT  →  BESOIN\nATELIER / STOCK  →  PRODUITS DISPONIBLES')
b.pic(KIT/'00_CARTE_GLOBALE/G01_carte_globale_3_flux.png','Trois flux, un même besoin client. Parcours principal proposé à valider.',14.5)
b.p('Opérateurs atelier · responsables atelier et stock · utilisateurs ERPNext métier')
b.box('Version de travail','Les captures de démonstration et les vues reconstituées ne sont pas des preuves de transactions exécutées. Aucune session métier authentifiée n’est présentée comme capturée.')
b.section('sommaire','Sommaire des parcours','Codes stables pour naviguer et commenter le document.')
b.table(['Partie','Chapitres','Question utilisateur'],[['Comprendre','00–05','Comment les trois flux se rejoignent-ils ?'],['I · Fournisseur','F01–F09','D’où vient la vitre et quand est-elle disponible ?'],['II · Client','C01–C13','Que demande le client et comment livrer / facturer ?'],['III · Atelier','A01–A28','Comment planifier, corriger et compenser ?'],['IV · Stock / Recomposition','S01–S14','Comment matérialiser les transformations ?'],['V · Exceptions','E01–E12','Que faire face à un écart ?'],['VI · Exercices / validation','X01–X03, V01','Comment tester et approuver le parcours ?'],['Glossaire / annexes / index','GLOSS, N01–N02, I01','Où trouver définitions, limites et captures ?']],'Carte du guide')
b.p('Lire F01–F09 pour la matière, C01–C07 pour la demande, puis A01–S14. Le Flux Client reprend en C08. Les opérations A12–A17 sont des choix selon le constat, pas une série d’actions toutes obligatoires.')
b.section('00','00 · Bienvenue dans mat_reco')
b.label('objectif','Objectif','Comprendre le parcours complet avant les écrans.')
b.p('mat_reco organise les matières dimensionnées, leurs pièces et leurs chutes. Le fournisseur fournit la matière ; le client exprime un besoin ; l’atelier rapproche les deux. Les Repack constatent les transformations en stock.')
b.table(['Préparer','Constater','Transformer'],[['Devis / commande et Planification MCP','Retour Terrain et incidents','Stock Entry Repack soumis'],['Besoin et plan à contrôler','Résultat effectif à enregistrer','Pièces / ouvrages réellement disponibles']],'Trois niveaux à distinguer')
b.box('Ne pas confondre','Un dessin, un nom de série prévu ou un brouillon ne prouvent pas une disponibilité physique en stock.')
b.label('suivante','Ensuite','01 — lire la carte globale ; puis 02 — comprendre les liens des documents.')
b.section('01','01 · Carte globale des trois flux')
b.pic(KIT/'00_CARTE_GLOBALE/G01_carte_globale_3_flux.png','Schéma G01 · Deux entrées convergent vers le MCP : matière disponible et besoin dimensionné.',23.0)
b.section('02','02 · Documents ERPNext et relations')
b.pic(KIT/'06_FLOWCHARTS/G06_relations_documents.png','Schéma G06 · Repack de découpe et Repack de recomposition sont deux opérations distinctes.',19.5)
b.label('suivante','Jonction','C07 → A01 transmet le besoin ; S14 → C08 rend le produit au parcours client.')
b.section('03','03 · Rôles et responsabilités')
b.table(['Rôle proposé','Responsabilité','À confirmer'],[['Achats / stock','Approvisionnement et réception','Choix PO / PR / Receipt Plan'],['Commercial','Devis et besoin confirmé','Création / soumission de commande'],['Opérateur atelier','Constats et opérations terrain','Droits sur incidents'],['Responsable atelier','Couverture et plan effectif','Droit de valider la compensation'],['Responsable stock','Contrôle et soumission Repack','Ordre avec la soumission MCP'],['Expédition / facturation','Livraison, facture, références','Chemins et traitement stock'],['Administrateur','Articles, scripts, droits','Les DocTypes livrent System Manager']],'Responsabilités à approuver')
b.box('Décision ouverte V01','L’ordre de soumission du MCP et des Repack ne constitue pas encore une procédure officielle.')
b.section('04','04 · Lire les repères visuels')
b.label('position','Position','Le mini-graphe marque la phase courante et le code de fiche.')
b.table(['Couleur','Usage'],[['Bleu','Préparation, atelier et action'],['Vert','Résultat, disponibilité et stock réel'],['Orange','Contrôle, attention et décision'],['Rouge','Blocage, risque ou écart'],['Violet / gris','Client, contexte et navigation']],'Code couleur')
for icon,label,txt in [('objectif','Objectif','Ce qu’il faut obtenir.'),('pourquoi','Pourquoi','L’utilité métier de l’étape.'),('entrees','Entrées','Informations nécessaires.'),('action','Action','Gestes à réaliser.'),('controle','Contrôle','Points à vérifier.'),('resultat','Résultat','Condition avant de poursuivre.'),('suivante','Étape suivante','Transition vers un chapitre.'),('attention','Attention','Portée et erreurs fréquentes.'),('astuce','Astuce','Lire les images en pleine résolution depuis l’index Excel.')]:b.label(icon,label,txt)
b.p('Démonstration : code réel de composant, données / réponses locales simulées. Reconstituée : illustration d’un formulaire à partir des champs et procédures. Réelle : capture de session ERPNext ; aucune n’est revendiquée dans ce lot.')
b.section('05','05 · Cas pratique fil rouge')
b.table(['Élément','Valeur fictive','Contrôle'],[['Matière','2 plaques 3210 × 2000 × 6 mm','12,84 m² au total'],['Commande SO-DEMO-001','2 OUVRAGE-DEMO','2 DECOUPE-6MM par ouvrage'],['Besoins atelier','4 pièces 800 × 600 × 6 mm','1,92 m² de produits découpés'],['Service ligne parent','PERIMETRE pour 2 ouvrages','2,8 m × 2 = 5,6 m'],['Scénario de casse','Pièce P2 détruite','3 conformes + 1 à remplacer'],['Stock complémentaire','Chute 600 × 800 × 6','Rotation terrain à 90° possible']],'Jeu cohérent de formation')
b.box('Simplifications explicites','Les scènes géométriques utilisent kerf = 0 pour montrer clairement les gestes. Utiliser le kerf réel dans la recette. La variante Chute en L utilise une scène distincte 1800 × 1500 mm.')
b.p('PO / PR et Material Receipt Plan sont deux scénarios alternatifs de réception des deux plaques : ne pas les additionner. Les identifiants présents dans les illustrations sont fictifs.')
for group in ['F','C','A','S']:
 gtitle,gfolder,gcolor=GROUPS[group];b.section('part_'+group,gtitle,{'F':'Approvisionner la matière','C':'De la demande à la facture, avec passage explicite par le Flux 3','A':'Cœur du guide : planifier, constater, corriger et compenser','S':'Prolongement du Flux 3 : matérialiser la transformation et retourner au client'}[group])
 b.pic(KIT/'06_FLOWCHARTS'/(GRAPH_MAP[group]+'.png'),'Schéma '+{'F':'G02','C':'G03','A':'G04','S':'G06'}[group]+' · Carte détaillée de ce parcours.',23)
 for code in [c for c in all_codes() if c[0]==group]:
  if code=='A15':b.section('graph_casse','G07 · Décision en cas de casse');b.pic(KIT/'06_FLOWCHARTS/G07_decision_casse.png','Schéma G07 · Incidents, couverture et reprise en cas de casse successive.',23)
  if code=='A20':b.section('graph_comp','G05 · Compensation : priorités et validation');b.pic(KIT/'06_FLOWCHARTS/G05_compensation.png','Schéma G05 · Cliquer Calculer déclenche la recherche prioritaire des sources.',23)
  st=STEPS[code];b.section(code,code+' · '+st['title'],st['evidence'])
  b.pic(KIT/'06_FLOWCHARTS'/('position_'+code+'.png'),'Position dans le flux · étape '+code,1.85)
  b.label('role','Rôle',st['role'],'violet');b.label('document','Document',st['document'],'violet')
  b.label('objectif','Objectif',st['objective']);b.label('pourquoi','Pourquoi cette étape ?',st['why']);b.label('position','Quand l’utiliser',st['when'])
  captures=BYCODE.get(code,BYCODE.get(REUSE.get(code,''),[]))
  for c in captures:
   b.pic(KIT/c['chemin'],f"Figure {c['figure']} · {c['id']} — {c['message']} [{c['statut']} ; à confirmer sur site]",max_h=5.1 if len(captures)>1 else 7.0)
  if not captures:
   g=GRAPH_MAP[group];b.pic(KIT/'06_FLOWCHARTS'/(g+'.png'),'Repère de décision du parcours · voir aussi le schéma détaillé de la partie.',max_h=6.8)
  b.label('entrees','Entrées / prérequis',st['inputs']);b.action_controls(st['actions'],st['controls'])
  b.box('Résultat attendu',st['result'],'green');b.label('suivante','Étape suivante',st['next'],'green')
  linked=[]
  for e,v in EXCEPTIONS.items():
   origin=v[0]
   ranges=re.findall(r'([A-Z])(\d{2})–(?:[A-Z])?(\d{2})',origin)
   if code in re.findall(r'[A-Z]\d{2}',origin) or any(code[0]==g and int(a)<=int(code[1:])<=int(z) for g,a,z in ranges):linked.append(e)
  st['exceptions']=linked
  text=st['attention'] or 'Si un contrôle est incohérent, ne pas poursuivre sans résolution.'
  if linked:text+=' Voir '+', '.join(linked)+'.'
  b.box('Attention / erreurs fréquentes',text)
# Exceptions are connected to originating chapters, not isolated general warnings.
b.section('part_E','Partie V · Gestion des exceptions','Revenir à la fiche d’origine après correction, sans contourner les contrôles.')
b.table(['Code','Situation','Étapes liées'],[[f'E{i:02d}',t,EXCEPTIONS[f'E{i:02d}'][0]] for i,t in enumerate(TITLES['E'],1)],'Index des exceptions')
for code,(origin,checks,action,nxt) in EXCEPTIONS.items():
 b.section(code,code+' · '+title(code),'Vous êtes dans une exception du parcours : '+origin)
 b.label('position','Survient à',origin);b.label('objectif','Objectif','Résoudre l’écart avant de poursuivre le flux.')
 for c in BYCODE.get(code,[]):b.pic(KIT/c['chemin'],f"Figure {c['figure']} · {c['id']} — {c['message']} [{c['statut']}]",7.4)
 b.label('controle','Vérifier',checks,'orange');b.label('action','Action',action)
 b.box('Résultat attendu','Le contrôle bloquant est compris et résolu, ou transmis au responsable.','green');b.label('suivante','Reprendre à',nxt,'green')
 b.box('Attention','Ne pas modifier les séries ou les données techniques pour faire disparaître le message. Conserver référence, message exact et capture pour le responsable.')
# Guided user acceptance, never represented as already-executed business tests.
b.section('part_X','Partie VI · Exercices et validation','Jeu de recette isolé ; aucun exercice ci-dessous n’a été soumis sur des documents métier pour ce guide.')
b.p('Les trois exercices relient les flux. Ils seront exécutés avec les rôles du site et un responsable de validation. Les soumissions et annulations doivent suivre le circuit approuvé.')
for ex in EXERCISES:
 b.section(ex['id'],ex['id']+' · '+ex['title'],ex['route']);b.label('objectif','Objectif',ex['goal']);b.label('entrees','Entrées',ex['inputs']);b.action_controls(ex['actions'],ex['checks']);b.label('suivante','Ensuite',ex['next'],'green');b.box('Preuve à conserver','Documents, séries, messages et captures de la recette réellement exécutée, avec date et validateur.')
b.section('V01','V01 · Validation et décisions ouvertes','La version v0.2 est une base éditoriale, pas une procédure approuvée.')
b.table(['Priorité','Sujet','Responsable'],[[id,topic,owner] for id,topic,state,owner,test in DECISIONS[:8]],'Décisions de flux et de droits')
b.p('La liste complète est fournie séparément dans Points-a-valider-mat_reco-v0.2.docx et dans le classeur d’index. Faire valider ensuite les exceptions, les coûts et les captures de session ERPNext.')
b.label('controle','Fin de poste','Besoin couvert ; incidents appliqués ; mouvements contrôlés ; séries disponibles ; exceptions transmises.','orange')
b.p('Validé par : ____________________  Fonction : ____________________');b.p('Date : ____________________  Version approuvée : ____________________')
b.section('GLOSS','Glossaire')
b.table(['Terme','Lecture utilisateur'],GLOSSARY,'Lexique métier')
b.section('N01','N01 · Annexes : réglages et circuits complémentaires')
b.label('document','Reconfiguration Settings','Réglages administrateur : kerf, seuils de conservation, rotation, split, réutilisation, stratégie, algorithme et heuristique.')
b.p('Certains paramètres concernent surtout le moteur historique. Ne pas supposer qu’ils pilotent tous les écrans. Le réglage de Planification ne bloque jamais la rotation à 90° en Retour Terrain.')
b.label('document','Material Issue Plan','Sortie matière annexe : Reason, séries qualifiées et contrôle avant soumission. Le plan crée / soumet Material Issue ; ce n’est pas la compensation d’une pièce du MCP.')
for c in BYCODE.get('E08',[]):b.pic(KIT/c['chemin'],f"Figure {c['figure']} · {c['id']} — Circuit annexe de sortie matière [{c['statut']}].",7.5)
b.box('Material Production Plan','Structure de formulaire présente, contrôleur sans orchestration active. Ne pas présenter génération d’arbre, contrôle stock ou production comme un parcours automatisé confirmé.')
b.section('N02','N02 · Sources, limites et contrôle documentaire')
b.p('La refonte suit les deux architectures fournies le 14 octobre 2026. Les règles propres au module sont rapprochées des contrôleurs, services, hooks, fixtures et composants Vue. Les transitions achat / vente proviennent des sources ERPNext locales.')
b.table(['Catégorie','Signification'],[['Déduit du code','Fonction identifiée dans les sources ; pas une recette complète en production.'],['Standard ERPNext','Parcours de document standard ; états, droits et options restent dépendants du site.'],['Procédure proposée','Contrôle physique, responsabilités ou transition métier à approuver.'],['Historique','Circuit Material Reconfiguration avec effets de soumission / annulation spécifiques.'],['Hors périmètre automatisé','Material Production Plan et fonctions non orchestrées dans le contrôleur actuel.']],'Niveaux de preuve')
b.p('Provenance des 60 usages indexés : 26 démonstrations de composants réels et 34 illustrations reconstituées, avec 54 fichiers image distincts mutualisés. Aucune capture d’une session métier authentifiée n’est revendiquée. Les réponses de sauvegarde montrées sont locales et simulées.')
b.box('À remplacer / confirmer','Formulaires ERPNext fournisseur / client / stock, apparence des hôtes de dialogue, scripts réellement actifs, droits et formats imprimés.')
b.p('Le manifeste de sources et le dossier de validation permettent de contrôler les assertions et les visuels. Le PDF de lecture est rendu depuis le même contenu que le Word ; il n’est pas une conversion effectuée par Word.')
b.section('I01','I01 · Index des captures et navigation des assets')
b.p('Le fichier INDEX_CAPTURES.xlsx est à la racine de MANUEL_GLOBAL_MAT_RECO. Il contient identifiant, fichier, flux, étape, écran, message pédagogique, statut, source et annotations. Ouvrir les PNG en pleine résolution pour la formation ou Skywork.')
for group in ['F','C','A','S','E']:
 rows=[c for c in CAPS if c['etape'][0]==group]
 if not rows:continue
 if group!='F':b.section('index_'+group,'I01 · Index — '+GROUPS[group][0])
 b.table(['Figure / ID','Étape / écran','Statut'],[[str(c['figure'])+' · '+c['id'],c['etape']+' · '+c['ecran'],c['statut']] for c in rows],GROUPS[group][0])
b.save()
# Separate deliverable D, explicitly reviewable.
v=Book('Points à valider — mat_reco v0.2','Points-a-valider-mat_reco-v0.2');v.section('validation_cover','Points à valider · mat_reco v0.2','Contrôleur documentaire, responsables atelier / stock et administrateur du site')
v.p('Décisions ouvertes et limites à résoudre avant diffusion officielle. Le manuel et ses captures conservent ces réserves. Aucune soumission métier n’a été effectuée pour fabriquer ce lot.')
for offset in range(0,len(DECISIONS),4):
 v.section('validation_'+str(offset),f'Décisions {offset+1} à {min(offset+4,len(DECISIONS))}')
 for id,topic,state,owner,test in DECISIONS[offset:offset+4]:
  v.label('controle',id+' · '+topic,state,'orange');v.label('role','Responsable',owner,'violet');v.label('action','Validation attendue',test);v.p('Décision / preuve : __________________________________________________')
v.section('missing','Captures manquantes et contrôles de site')
v.table(['Groupe','Visuels de session réelle à obtenir'],[['Fournisseur','Purchase Order, Purchase Receipt, Material Receipt Plan, Serial No, stock disponible.'],['Client','Quotation, services / accessoires, Sales Order, Delivery Note, bundles, Sales Invoice.'],['Atelier','MCP complet, actions et compensation reliés à un compte habilité ; sauvegardes réellement confirmées.'],['Stock','Repack découpe, recomposition, mouvements soumis, coûts et nouveaux Serial No.'],['Historique / annexe','Material Reconfiguration, Material Issue Plan ; activation réelle des boutons.']],'Remplacement documentaire')
v.p('Les images de démonstration / reconstitution sont utilisables pour la base éditoriale mais ne doivent pas être renommées « capture réelle ERPNext » lors de la finalisation graphique.')
v.section('changes','Écarts avec la version v0.1')
v.list(['Passage de fonctions isolées à trois flux convergents.','64 fiches métier avec position, pourquoi, actions, contrôles, résultat et suite.','Repack de découpe séparé du Repack de recomposition.','Fournisseur et Client développés, avec transitions C07 → A01 et S14 → C08.','7 cartes de flux, 64 mini-graphes, 13 icônes et 60 captures indexées.','Index Excel, storyboard, sources éditoriales, exercices et décisions séparées.','Traçabilité de chaque visuel ; limites de session et de recette clairement conservées.'])
v.save()
(OUT/'edition/Fiches-metier-v0.2.json').write_text(json.dumps(STEPS,ensure_ascii=False,indent=2))
(OUT/'edition/Index-captures-v0.2.json').write_text(json.dumps(CAPS,ensure_ascii=False,indent=2))
(OUT/'edition/Plan-du-manuel-v0.2.json').write_text(json.dumps(b.sections,ensure_ascii=False,indent=2))
print('Phase 4 : manuel global, 64 fiches et document séparé de validation générés ;',len(b.sections),'sections de lecture.')
